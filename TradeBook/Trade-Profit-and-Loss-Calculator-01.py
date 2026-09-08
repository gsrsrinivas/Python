"""
Trade P&L Calculator -- Zerodha & Fyers
Drop all trade Excel files into the same folder as this script and run:
    python trade_pnl_calculator.py
Output: trade_pnl_output.xlsx

Matching rules (FIFO):
  - Buys are matched to sells chronologically per symbol.
  - EQUITY (Segment=EQ / Equity): sell date >= buy date (cannot be before buy date).
  - OPTIONS & FUTURES (Segment=Derivatives/FO/NFO etc.): no sell-date restriction
    (intraday & same-day squaring off are valid).
  - Partial quantities are handled: one buy lot can match multiple sells and vice-versa.
"""

import os
import sys
import glob
import warnings
from collections import deque
from datetime import datetime, date

import pandas as pd

warnings.filterwarnings("ignore")

# ---------------------------------------------------------------------------
# 1. Broker detection
# ---------------------------------------------------------------------------

ZERODHA_COLS = {
    "Account", "Period", "Symbol", "ISIN", "Trade Date",
    "Exchange", "Segment", "Series", "Trade Type", "Auction",
    "Quantity", "Price", "Trade ID", "Order ID", "Order Execution Time",
}

# Fyers column name uses literal ampersand: "Date & time"
FYERS_REQUIRED = {"Account", "Period", "Name", "Side", "Product type",
                  "Qty", "Traded price", "Total value", "Segment",
                  "Exchange order ID", "OMS order ID"}


def detect_broker(df: pd.DataFrame):
    cols = set(df.columns)
    if ZERODHA_COLS.issubset(cols):
        return "Zerodha"
    # Fyers: check required cols + a column containing both 'date' and 'time' words
    fyers_date_col = [c for c in cols if "date" in c.lower() and "time" in c.lower()]
    if FYERS_REQUIRED.issubset(cols) and fyers_date_col:
        return "Fyers"
    return None


def get_fyers_date_col(df: pd.DataFrame) -> str:
    """Return the actual date-time column name from a Fyers file."""
    for c in df.columns:
        if "date" in c.lower() and "time" in c.lower():
            return c
    # Fallback to the known literal name
    return "Date & time"


# ---------------------------------------------------------------------------
# 2. Date parsing
# ---------------------------------------------------------------------------

def parse_date_flexible(val):
    if val is None:
        return None
    if isinstance(val, float) and val != val:   # NaN
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    try:
        if isinstance(val, pd.Timestamp):
            return val.date()
    except Exception:
        pass
    s = str(val).strip()
    for fmt in (
        "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y",
        "%d %b %Y", "%d-%b-%Y",
        "%d %b %Y, %I:%M:%S %p", "%d %B %Y, %I:%M:%S %p",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    try:
        return pd.to_datetime(s, dayfirst=True).date()
    except Exception:
        return None


# ---------------------------------------------------------------------------
# 3. Normalise to unified schema
# ---------------------------------------------------------------------------

def normalise_zerodha(df: pd.DataFrame, source: str) -> pd.DataFrame:
    out = pd.DataFrame()
    out["broker"]               = "Zerodha"
    out["account"]              = df["Account"].astype(str).str.strip()
    out["symbol"]               = df["Symbol"].astype(str).str.strip().str.upper()
    out["isin"]                 = df["ISIN"].astype(str).str.strip()
    out["trade_date"]           = df["Trade Date"].apply(parse_date_flexible)
    out["side"]                 = df["Trade Type"].astype(str).str.strip().str.upper()
    out["quantity"]             = pd.to_numeric(df["Quantity"], errors="coerce").fillna(0).astype(int)
    out["price"]                = pd.to_numeric(df["Price"], errors="coerce").fillna(0.0)
    out["segment"]              = df["Segment"].astype(str).str.strip().str.upper()
    out["exchange"]             = df["Exchange"].astype(str).str.strip().str.upper()
    out["order_execution_time"] = df["Order Execution Time"].astype(str).str.strip()
    out["source_file"]          = os.path.basename(source)
    # Zerodha Segment column values: EQ = equity, FO/NFO/BFO/CDS = derivatives
    seg_upper = df["Segment"].astype(str).str.strip().str.upper()
    out["product_type"]         = seg_upper.apply(
        lambda s: "Equity" if s == "EQ" else "Derivatives"
    )
    return out


def normalise_fyers(df: pd.DataFrame, source: str) -> pd.DataFrame:
    date_col = get_fyers_date_col(df)
    out = pd.DataFrame()
    out["broker"]               = "Fyers"
    out["account"]              = df["Account"].astype(str).str.strip()
    out["symbol"]               = df["Name"].astype(str).str.strip().str.upper()
    out["isin"]                 = ""
    out["trade_date"]           = df[date_col].apply(parse_date_flexible)
    out["side"]                 = df["Side"].astype(str).str.strip().str.upper()
    out["quantity"]             = pd.to_numeric(df["Qty"], errors="coerce").fillna(0).astype(int)
    out["price"]                = pd.to_numeric(df["Traded price"], errors="coerce").fillna(0.0)
    out["segment"]              = df["Segment"].astype(str).str.strip().str.upper()
    out["exchange"]             = ""
    out["order_execution_time"] = df[date_col].astype(str).str.strip()
    out["source_file"]          = os.path.basename(source)
    # Fyers Segment column: "Equity" = equity, "Derivatives" / "Commodity" = F&O/Commodity
    seg_upper = df["Segment"].astype(str).str.strip().str.upper()
    out["product_type"]         = seg_upper.apply(
        lambda s: "Equity" if s == "EQUITY" else "Derivatives"
    )
    return out


# ---------------------------------------------------------------------------
# 4. Load all .xlsx files from folder
# ---------------------------------------------------------------------------

def load_folder(folder: str) -> pd.DataFrame:
    # Use os.listdir so dot-prefixed files (e.g. .TradeBook...xlsx) are found too
    all_names = os.listdir(folder)
    files = sorted(
        os.path.join(folder, n)
        for n in all_names
        if n.lower().endswith(".xlsx")
    )
    if not files:
        print("ERROR: No .xlsx files found in: " + folder)
        sys.exit(1)

    frames = []
    for fpath in files:
        fname = os.path.basename(fpath)
        # Skip the output file itself
        if fname == "trade_pnl_output.xlsx":
            continue
        print("  Loading: " + fname)
        try:
            raw = pd.read_excel(fpath)
        except Exception as exc:
            print("    Skipped (read error): " + str(exc))
            continue
        broker = detect_broker(raw)
        if broker is None:
            print("    Skipped (unrecognised format)")
            continue
        print("    Detected: " + broker + "  (" + str(len(raw)) + " rows)")
        if broker == "Zerodha":
            frames.append(normalise_zerodha(raw, fpath))
        elif broker == "Fyers":
            frames.append(normalise_fyers(raw, fpath))

    if not frames:
        print("ERROR: No valid trade files were loaded.")
        sys.exit(1)

    combined = pd.concat(frames, ignore_index=True)
    combined.dropna(subset=["trade_date"], inplace=True)
    combined.sort_values(["symbol", "trade_date", "order_execution_time"], inplace=True)
    combined.reset_index(drop=True, inplace=True)
    return combined


# ---------------------------------------------------------------------------
# 5. FIFO matching engine
# ---------------------------------------------------------------------------


# ---------------------------------------------------------------------------
# Helper: decide whether a trade row is equity (vs derivatives/F&O)
# ---------------------------------------------------------------------------

def is_equity(rd: dict) -> bool:
    """Return True if this trade row is an equity share (not F&O / commodity)."""
    pt = str(rd.get("product_type", "")).upper()
    return pt == "EQUITY"


def fifo_match(trades: pd.DataFrame):
    matched_rows = []
    open_rows    = []

    for symbol, grp in trades.groupby("symbol", sort=False):
        grp = grp.sort_values("trade_date").reset_index(drop=True)
        buy_queue: deque = deque()

        for _, row in grp.iterrows():
            rd = row.to_dict()
            side = rd["side"]

            if side == "BUY":
                buy_queue.append({"qty_remaining": int(rd["quantity"]), "data": rd})

            elif side == "SELL":
                srq = int(rd["quantity"])

                while srq > 0:
                    # Case 1: No matching buy exists
                    if not buy_queue:
                        open_rows.append({
                            "Type":         "Unmatched SELL",
                            "Broker":       rd["broker"],
                            "Account":      rd["account"],
                            "Symbol":       symbol,
                            "ISIN":         rd["isin"],
                            "Product Type": rd.get("product_type", ""),
                            "Buy Date":     None,
                            "Buy Price":    None,
                            "Sell Date":    rd["trade_date"],
                            "Sell Price":   rd["price"],
                            "Quantity":     srq,
                            "Buy Value":    None,
                            "Sell Value":   None,
                            "PnL":          None,
                            "PnL %":        None,
                            "Holding Days": None,
                            "Source File":  rd["source_file"],
                            "Note":         "No matching BUY found (possible short-sell or missing buy leg)",
                        })
                        srq = 0
                        break

                    buy = buy_queue[0]
                    bd  = buy["data"]["trade_date"]

                    # Case 2: Date-rule check (equity only)
                    # Equity: sell date must be >= buy date (cannot be before buy date)
                    # Derivatives / F&O: no date restriction
                    if is_equity(rd) and rd["trade_date"] < bd:
                        open_rows.append({
                            "Type":         "Rule Violation",
                            "Broker":       rd["broker"],
                            "Account":      rd["account"],
                            "Symbol":       symbol,
                            "ISIN":         rd["isin"],
                            "Product Type": rd.get("product_type", ""),
                            "Buy Date":     bd,
                            "Buy Price":    buy["data"]["price"],
                            "Sell Date":    rd["trade_date"],
                            "Sell Price":   rd["price"],
                            "Quantity":     srq,
                            "Buy Value":    None,
                            "Sell Value":   None,
                            "PnL":          None,
                            "PnL %":        None,
                            "Holding Days": None,
                            "Source File":  rd["source_file"],
                            "Note":         "SELL date is BEFORE BUY date (equity rule violation, skipped)",
                        })
                        srq = 0
                        break

                    # Case 3: Normal FIFO match
                    mq  = min(buy["qty_remaining"], srq)
                    bv  = mq * buy["data"]["price"]
                    sv  = mq * rd["price"]
                    pnl = sv - bv
                    pp  = (pnl / bv * 100) if bv else None
                    hd  = (rd["trade_date"] - bd).days

                    matched_rows.append({
                        "Broker":        rd["broker"],
                        "Account":       rd["account"],
                        "Symbol":        symbol,
                        "ISIN":          rd["isin"],
                        "Product Type":  rd.get("product_type", ""),
                        "Buy Date":      bd,
                        "Buy Price":     round(buy["data"]["price"], 4),
                        "Sell Date":     rd["trade_date"],
                        "Sell Price":    round(rd["price"], 4),
                        "Quantity":      mq,
                        "Buy Value":     round(bv, 2),
                        "Sell Value":    round(sv, 2),
                        "PnL":           round(pnl, 2),
                        "PnL %":         round(pp, 4) if pp is not None else None,
                        "Holding Days":  hd,
                        "Trade Type":    "Long-Term" if hd >= 365 else "Short-Term",
                        "Source File":   rd["source_file"],
                    })

                    buy["qty_remaining"] -= mq
                    srq                 -= mq
                    if buy["qty_remaining"] == 0:
                        buy_queue.popleft()

        # Remaining buys = open positions
        for buy in buy_queue:
            if buy["qty_remaining"] > 0:
                open_rows.append({
                    "Type":         "Open Position (BUY held)",
                    "Broker":       buy["data"]["broker"],
                    "Account":      buy["data"]["account"],
                    "Symbol":       symbol,
                    "ISIN":         buy["data"]["isin"],
                    "Product Type": buy["data"].get("product_type", ""),
                    "Buy Date":     buy["data"]["trade_date"],
                    "Buy Price":    buy["data"]["price"],
                    "Sell Date":    None,
                    "Sell Price":   None,
                    "Quantity":     buy["qty_remaining"],
                    "Buy Value":    round(buy["qty_remaining"] * buy["data"]["price"], 2),
                    "Sell Value":   None,
                    "PnL":          None,
                    "PnL %":        None,
                    "Holding Days": None,
                    "Source File":  buy["data"]["source_file"],
                    "Note":         "Position not yet closed",
                })

    return pd.DataFrame(matched_rows), pd.DataFrame(open_rows)


# ---------------------------------------------------------------------------
# 6. Summary builders
# ---------------------------------------------------------------------------

def build_symbol_summary(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return pd.DataFrame()
    g = m.groupby(["Symbol", "Trade Type"]).agg(
        Total_Buy_Value  =("Buy Value",  "sum"),
        Total_Sell_Value =("Sell Value", "sum"),
        Total_PnL        =("PnL",        "sum"),
        Num_Trades       =("Quantity",   "count"),
        Total_Qty        =("Quantity",   "sum"),
    ).reset_index()
    g["Avg_PnL_Pct"] = ((g["Total_Sell_Value"] - g["Total_Buy_Value"])
                        / g["Total_Buy_Value"] * 100).round(2)
    return g.sort_values(["Symbol", "Trade Type"]).reset_index(drop=True)


def build_overall(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return pd.DataFrame()
    st_pnl = m.loc[m["Trade Type"] == "Short-Term", "PnL"].sum()
    lt_pnl = m.loc[m["Trade Type"] == "Long-Term",  "PnL"].sum()
    return pd.DataFrame({
        "Metric": [
            "Total Realised PnL (INR)",
            "Total Buy Value (INR)",
            "Total Sell Value (INR)",
            "Matched Trade Lots",
            "Profitable Trade Lots",
            "Loss Trade Lots",
            "Short-Term Capital Gains (INR)",
            "Long-Term Capital Gains (INR)",
        ],
        "Value": [
            round(m["PnL"].sum(), 2),
            round(m["Buy Value"].sum(), 2),
            round(m["Sell Value"].sum(), 2),
            len(m),
            int((m["PnL"] > 0).sum()),
            int((m["PnL"] < 0).sum()),
            round(st_pnl, 2),
            round(lt_pnl, 2),
        ],
    })


def build_monthly(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return pd.DataFrame()
    df = m.copy()
    df["Month"] = pd.to_datetime(df["Sell Date"]).dt.to_period("M").astype(str)
    r = df.groupby("Month").agg(
        Buy_Value  =("Buy Value",  "sum"),
        Sell_Value =("Sell Value", "sum"),
        PnL        =("PnL",        "sum"),
        Num_Trades =("Quantity",   "count"),
    ).reset_index()
    r["PnL_Pct"] = ((r["Sell_Value"] - r["Buy_Value"]) / r["Buy_Value"] * 100).round(2)
    return r.sort_values("Month").reset_index(drop=True)


def build_broker_summary(m: pd.DataFrame) -> pd.DataFrame:
    if m.empty:
        return pd.DataFrame()
    b = m.groupby("Broker").agg(
        Total_Buy_Value  =("Buy Value",  "sum"),
        Total_Sell_Value =("Sell Value", "sum"),
        Total_PnL        =("PnL",        "sum"),
        Matched_Trades   =("Quantity",   "count"),
    ).reset_index()
    b["PnL_Pct"] = ((b["Total_Sell_Value"] - b["Total_Buy_Value"])
                    / b["Total_Buy_Value"] * 100).round(2)
    return b


# ---------------------------------------------------------------------------
# 7. Excel styling
# ---------------------------------------------------------------------------

def _style_ws(ws, df: pd.DataFrame) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    HDR_FILL = PatternFill("solid", fgColor="1F4E79")
    ALT_FILL = PatternFill("solid", fgColor="D6E4F0")
    GRN_FILL = PatternFill("solid", fgColor="C6EFCE")
    RED_FILL = PatternFill("solid", fgColor="FFC7CE")
    THIN     = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"),  bottom=Side(style="thin"),
    )

    def is_money(n):
        n = n.lower()
        return any(k in n for k in ["value", "price", "pnl", "total_pnl",
                                     "total_buy", "total_sell", "gains"])
    def is_pct(n):   return "pct" in n.lower() or n.endswith("%")
    def is_date(n):  return "date" in n.lower()
    def is_int(n):
        return n.lower() in ("quantity", "holding days", "num_trades",
                              "matched_trades", "total_qty")

    # Find PnL column index (1-based)
    pnl_idx = next(
        (i + 1 for i, c in enumerate(df.columns) if c.lower() in ("pnl", "total_pnl")),
        None,
    )

    # Header row
    for cell in ws[1]:
        cell.font      = Font(bold=True, color="FFFFFF", size=10)
        cell.fill      = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border    = THIN

    # Data rows
    for r_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for c_idx, cell in enumerate(row, start=1):
            cn = df.columns[c_idx - 1]
            cell.border    = THIN
            cell.alignment = Alignment(vertical="center")

            if is_money(cn) and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0.00"
            elif is_pct(cn) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            elif is_date(cn):
                cell.number_format = "DD-MMM-YYYY"
                cell.alignment = Alignment(horizontal="center")
            elif is_int(cn) and isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

            # PnL colour coding
            if pnl_idx and c_idx == pnl_idx and isinstance(cell.value, (int, float)):
                cell.fill = GRN_FILL if cell.value >= 0 else RED_FILL
            elif r_idx % 2 == 0:
                cell.fill = ALT_FILL

    # Auto column width
    for col in ws.iter_cols():
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 42)

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def write_excel(path: str, matched, open_pos, symbol_summ, monthly, broker_s, overall) -> None:
    from openpyxl import load_workbook

    sheets = [
        ("Overall Summary", overall),
        ("Symbol Summary",  symbol_summ),
        ("Monthly PnL",     monthly),
        ("Broker Summary",  broker_s),
        ("Matched Trades",  matched),
        ("Open & Unmatched", open_pos),
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets:
            if df is not None and not df.empty:
                df.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(path)
    for name, df in sheets:
        if df is not None and not df.empty and name in wb.sheetnames:
            _style_ws(wb[name], df)
    wb.save(path)
    print("  Saved: " + path)


# ---------------------------------------------------------------------------
# 8. Main
# ---------------------------------------------------------------------------

def main():
    folder = os.path.dirname(os.path.abspath(__file__))
    output = os.path.join(folder, "trade_pnl_output.xlsx")

    sep = "=" * 60
    print(sep)
    print("  Trade P&L Calculator  --  Zerodha & Fyers")
    print(sep)
    print("Scanning: " + folder)

    trades = load_folder(folder)
    print("Rows loaded  : " + str(len(trades)))
    print("Unique symbols: " + str(trades["symbol"].nunique()))
    print("Date range   : " + str(trades["trade_date"].min()) + " to " + str(trades["trade_date"].max()))

    print("Running FIFO matching ...")
    matched, open_pos = fifo_match(trades)
    print("Matched lots : " + str(len(matched)))
    print("Open/Unmatched: " + str(len(open_pos)))

    if not matched.empty:
        st = matched.loc[matched["Trade Type"] == "Short-Term", "PnL"].sum()
        lt = matched.loc[matched["Trade Type"] == "Long-Term",  "PnL"].sum()
        total = matched["PnL"].sum()
        print("Total Realised PnL  : INR " + "{:,.2f}".format(total))
        print("Short-Term PnL      : INR " + "{:,.2f}".format(st))
        print("Long-Term  PnL      : INR " + "{:,.2f}".format(lt))

    print("Writing Excel ...")
    write_excel(
        output, matched, open_pos,
        build_symbol_summary(matched),
        build_monthly(matched),
        build_broker_summary(matched),
        build_overall(matched),
    )

    print(sep)
    print("  Done! Open trade_pnl_output.xlsx to view results.")
    print(sep)


if __name__ == "__main__":
    main()
